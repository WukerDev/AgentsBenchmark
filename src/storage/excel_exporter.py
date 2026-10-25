from __future__ import annotations
from pathlib import Path
from typing import Dict, Any, List
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelBenchmarkExporter:
    """
    Exports comprehensive multi-tab formatted Excel reports for research analysis.
    """
    @classmethod
    def export(cls, output_path: Path, df_results: pd.DataFrame, stats: Dict[str, Any]):
        output_path.parent.mkdir(parents=True, exist_ok=True)

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            # Sheet 1: Main Task Results
            df_results.to_excel(writer, sheet_name="Wyniki_Zadan", index=False)

            # Sheet 2: Category Breakdown
            if "category" in df_results:
                cat_summary = df_results.groupby("category").agg(
                    Liczba_Zadan=("task_id", "count"),
                    Solo_Średnia=("score_solo", "mean"),
                    Multi_Średnia=("score_multi_final", "mean"),
                    Solo_Czas_s=("solo_time", "mean"),
                    Multi_Czas_s=("multi_time", "mean"),
                    Solo_Koszt_USD=("solo_cost_usd", "mean"),
                    Multi_Koszt_USD=("multi_cost_usd", "mean")
                ).round(2).reset_index()
                cat_summary.to_excel(writer, sheet_name="Podsumowanie_Kategorii", index=False)

            # Sheet 3: Turn Evolution
            if "score_multi_5" in df_results and "score_multi_10" in df_results:
                evo_df = df_results[["task_id", "score_solo", "score_multi_5", "score_multi_10", "score_multi_final"]].copy()
                evo_df.rename(columns={
                    "score_solo": "Solo_ReAct",
                    "score_multi_5": "Multi_Tura_5",
                    "score_multi_10": "Multi_Tura_10",
                    "score_multi_final": "Multi_Finał"
                }, inplace=True)
                evo_df.to_excel(writer, sheet_name="Ewolucja_Tur", index=False)

            # Sheet 4: Hypothesis Testing & Global Stats
            if stats:
                stat_rows = [
                    {"Metryka": "Liczba zbadanych zadań", "Wartość": stats.get("n_tasks")},
                    {"Metryka": "Średni wynik Solo (ReAct)", "Wartość": stats.get("solo_mean")},
                    {"Metryka": "Odchylenie standardowe Solo", "Wartość": stats.get("solo_std")},
                    {"Metryka": "Mediana Solo", "Wartość": stats.get("solo_median")},
                    {"Metryka": "Średni wynik Multi-Agent", "Wartość": stats.get("multi_mean")},
                    {"Metryka": "Odchylenie standardowe Multi-Agent", "Wartość": stats.get("multi_std")},
                    {"Metryka": "Mediana Multi-Agent", "Wartość": stats.get("multi_median")},
                    {"Metryka": "Różnica średnich (Delta)", "Wartość": stats.get("score_delta")},
                    {"Metryka": "Zwycięstwa Multi-Agent", "Wartość": stats.get("multi_wins")},
                    {"Metryka": "Remisy", "Wartość": stats.get("ties")},
                    {"Metryka": "Porażki Multi-Agent", "Wartość": stats.get("multi_losses")},
                    {"Metryka": "Win Rate Multi-Agent (%)", "Wartość": stats.get("multi_win_rate")},
                    {"Metryka": "Test t-Studenta (p-value)", "Wartość": stats.get("t_test_pvalue")},
                    {"Metryka": "Test t Istotny statystycznie (p < 0.05)", "Wartość": stats.get("t_test_significant")},
                    {"Metryka": "Test Wilcoxona (p-value)", "Wartość": stats.get("wilcoxon_pvalue")},
                    {"Metryka": "Wilcoxon Istotny statystycznie (p < 0.05)", "Wartość": stats.get("wilcoxon_significant")}
                ]
                pd.DataFrame(stat_rows).to_excel(writer, sheet_name="Statystyka_Naukowa", index=False)

            # Format Workbook Styles
            wb = writer.book
            header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True, name="Segoe UI")
            cell_font = Font(name="Segoe UI", size=10)

            for ws in wb.worksheets:
                ws.views.sheetView[0].showGridLines = True
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.font = cell_font

                for col in ws.columns:
                    col_letter = get_column_letter(col[0].column)
                    max_len = max(len(str(c.value or '')) for c in col)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
